import logging
from pptx import Presentation
from typing import List, Dict, Any
import os

logger = logging.getLogger(__name__)

def generate_presentation(title: str, content_sections: List[Dict[str, Any]], output_path: str = "output/presentation.pptx") -> str:
    """
    Generates a PowerPoint presentation from structured content sections.

    Args:
        title: The title of the presentation.
        content_sections: A list of dictionaries, where each dictionary
                           represents a section of the presentation. Each must have
                           'title' (str) and 'content' (list of strings).
        output_path: The file path to save the generated PPTX.

    Returns:
        The absolute path to the generated presentation file, or None on failure.
    """
    try:
        prs = Presentation()

        # Slide layout indices (Assuming standard layouts)
        # 0: Title Slide
        # 1: Title and Content
        # 5: Title Only

        # 1. Title Slide (Title)
        title_slide_layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(title_slide_layout)
        title_shape = slide.shapes.title
        subtitle_shape = slide.placeholders[1]
        title_shape.text = title
        subtitle_shape.text = "Workflow Completion Report"

        # 2. Content Slides
        for i, section in enumerate(content_sections):
            section_title = section.get("title", f"Section {i+1}")
            content_list = section.get("content", [])

            # Use Title and Content layout
            title_content_layout = prs.slide_layouts[1]
            slide = prs.slides.add_slide(title_content_layout)

            # Set the slide title
            slide.shapes.title.text = section_title

            # Add content bullet points
            body = slide.shapes.placeholders[1]
            tf = body.text_frame
            tf.clear() # Clear any default text

            for content_item in content_list:
                if isinstance(content_item, str) and content_item:
                    p = tf.add_paragraph()
                    p.text = content_item
                    p.level = 0 # Top level bullet point

        # Ensure output directory exists
        os.makedirs(os.path.dirname(output_path), exist_ok=True)

        # Save the presentation
        prs.save(output_path)
        output_path_abs = os.path.abspath(output_path)
        logging.info(f"Successfully generated presentation at: {output_path_abs}")
        return output_path_abs

    except Exception as e:
        logging.error(f"Error generating presentation: {e}")
        return None